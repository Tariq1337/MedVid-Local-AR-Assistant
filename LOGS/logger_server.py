import os
import re
import json
import time
import shutil
import base64
import soundfile as sf
from datetime import datetime
from fastapi import FastAPI, UploadFile, File, Form, Request
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
import uvicorn
from openai import OpenAI

# EXCEL LIBRARY
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# LOCAL AI LIBRARIES
from faster_whisper import WhisperModel
from kokoro_onnx import Kokoro
from search_engine import MedVidSearch

# ==============================================================================
#  CONFIGURATION
# ==============================================================================
LOG_FOLDER = "Mission_Logs"
EXCEL_FILE = "mission_log.xlsx"
VIDEO_DIR = r"C:\Users\tbahaaal\Documents\LOGS\MedVid_DATA\videos"
TRANSCRIPTS_FILE = "MedVid_DATA/transcripts.json"

# 1. LOCAL BRAIN (RTX 4090 - vLLM)
LOCAL_LLM_URL = "http://localhost:22002/v1"
LOCAL_MODEL_NAME = "Qwen/Qwen3-VL-8B-Instruct"

# 2. LOCAL EARS (Faster-Whisper)
print("👂 Loading Whisper (Ears)...")
whisper_model = WhisperModel("large-v3", device="cuda", compute_type="float16")

# 3. LOCAL MOUTH (Kokoro TTS)
print("👄 Loading Kokoro (Mouth)...")
kokoro = Kokoro("Kokoro/kokoro-v0_19.onnx", "Kokoro/voices.json")
VOICE_NAME = "af_bella"

# 4. MEDICAL SEARCH ENGINE
print("🔍 Loading Medical Search Engine...")
search_engine = MedVidSearch()

# 5. LLM CLIENT (reusable)
llm_client = OpenAI(base_url=LOCAL_LLM_URL, api_key="EMPTY")

# 6. VIDEO TRANSCRIPTS (for summarizing videos)
print("📜 Loading Video Transcripts...")
video_transcripts = {}
try:
    with open(TRANSCRIPTS_FILE, "r", encoding="utf-8") as f:
        video_transcripts = json.load(f)
    print(f"   ✅ Loaded transcripts for {len(video_transcripts)} videos")
except Exception as e:
    print(f"   ⚠️ Could not load transcripts: {e}")

# 7. YOLO CLASSES (for detection keyword scan)
YOLO_CLASSES = [
    "person", "bicycle", "car", "motorcycle", "airplane", "bus", "train", "truck", "boat", "traffic light",
    "fire hydrant", "stop sign", "parking meter", "bench", "bird", "cat", "dog", "horse", "sheep", "cow",
    "elephant", "bear", "zebra", "giraffe", "backpack", "umbrella", "handbag", "tie", "suitcase", "frisbee",
    "skis", "snowboard", "sports ball", "kite", "baseball bat", "baseball glove", "skateboard", "surfboard",
    "tennis racket", "bottle", "wine glass", "cup", "fork", "knife", "spoon", "bowl", "banana", "apple",
    "sandwich", "orange", "broccoli", "carrot", "hot dog", "pizza", "donut", "cake", "chair", "couch",
    "potted plant", "bed", "dining table", "toilet", "tv", "laptop", "mouse", "remote", "keyboard", "cell phone",
    "microwave", "oven", "toaster", "sink", "refrigerator", "book", "clock", "vase", "scissors", "teddy bear",
    "hair drier", "toothbrush", "monitor"
]

# ==============================================================================
#  FASTAPI APP
# ==============================================================================
app = FastAPI(title="MedVid AR Assistant")

app.mount("/videos", StaticFiles(directory="MedVid_DATA/videos"), name="videos")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=[
        "X-Detected-Object", "X-User-Transcript", "X-AI-Response",
        "X-Video-URL", "X-Video-Start", "X-Video-End", "X-Video-Title",
        "Access-Control-Expose-Headers"
    ]
)

# ==============================================================================
#  EXCEL ENGINE (ASU Maroon & Gold)
# ==============================================================================
ASU_MAROON = "8C1D40"
ASU_GOLD = "FFC627"

def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Mission Log"
        headers = [
            "Date", "Time", "Image", "Video Used",
            "Detected Items", "User Question", "AI Assessment", "AI Answer",
            "User Audio", "AI Audio"
        ]
        ws.append(headers)
        column_widths = [14, 12, 18, 30, 30, 40, 35, 55, 18, 18]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        header_fill = PatternFill(start_color=ASU_MAROON, end_color=ASU_MAROON, fill_type="solid")
        header_font = Font(bold=True, color=ASU_GOLD, size=12)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        wb.save(EXCEL_FILE)
        print(f"   📊 Created {EXCEL_FILE} (ASU Maroon & Gold)")

# ==============================================================================
#  MEMORY SYSTEM (includes video info)
# ==============================================================================
def get_full_memory(user_query=""):
    if not os.path.exists(EXCEL_FILE):
        return ""
    memory_context = ""
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        all_rows = list(ws.iter_rows(min_row=2, values_only=True))
        wb.close()
        if not all_rows:
            return ""
        # Tier 1: Last 10
        recent_rows = all_rows[-10:]
        memory_context += "RECENT CONVERSATIONS (last 10):\n"
        for row in recent_rows:
            if row and len(row) >= 8:
                q = row[5]   # User Question
                a = row[7]   # AI Answer
                v = row[3]   # Video Used
                if q and a:
                    memory_context += f"- User: '{q}'\n  AI: '{a}'\n"
                    if v and v != "N/A":
                        memory_context += f"  [VIDEO PLAYED: {v}]\n"
        # Tier 2: Deep scan if recall keywords detected
        recall_keywords = [
            "remember", "recall", "earlier", "before", "last time",
            "my name", "i said", "i told", "i asked", "i mentioned",
            "did i", "what did", "who am i", "do you know me",
            "previous", "history", "forgot", "forget"
        ]
        query_lower = user_query.lower()
        needs_deep_scan = any(kw in query_lower for kw in recall_keywords)
        if needs_deep_scan and len(all_rows) > 10:
            memory_context += "\nFULL CONVERSATION HISTORY (deep scan):\n"
            for row in all_rows[:-10]:
                if row and len(row) >= 8:
                    q, a, v = row[5], row[7], row[3]
                    if q and a:
                        memory_context += f"- User: '{q}'\n  AI: '{a}'\n"
                        if v and v != "N/A":
                            memory_context += f"  [VIDEO PLAYED: {v}]\n"
            print("   🧠 Deep memory scan triggered")
    except Exception as e:
        print(f"   ⚠️ Memory error: {e}")
    return memory_context

def get_last_video_id():
    """Find the most recently played video ID from the Excel log."""
    if not os.path.exists(EXCEL_FILE):
        return None
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        wb.close()
        for row in reversed(rows):
            if row and len(row) >= 4 and row[3] and row[3] != "N/A":
                # Extract video ID from "filename.mp4 [300s-368s]"
                video_str = str(row[3])
                if ".mp4" in video_str:
                    vid_id = video_str.split(".mp4")[0].strip()
                    return vid_id
    except Exception:
        pass
    return None

# ==============================================================================
#  EXCEL LOGGING
# ==============================================================================
def append_to_log(question, assessment, answer, img_path, video_used, detected_items, audio_in_path, audio_out_path):
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active
    now = datetime.now()
    row_data = [
        now.strftime("%Y-%m-%d"),
        now.strftime("%H:%M:%S"),
        "View Image",
        video_used if video_used else "N/A",
        detected_items if detected_items else "none",
        question,
        assessment,
        answer,
        "Play User",
        "Play AI"
    ]
    ws.append(row_data)
    new_row = ws[ws.max_row]
    center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    row_num = ws.max_row
    if row_num % 2 == 0:
        row_fill = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")
    else:
        row_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    for cell in new_row:
        cell.alignment = center_wrap
        cell.border = thin_border
        cell.fill = row_fill
    # AI Assessment and AI Answer columns: left-aligned for readability
    new_row[6].alignment = left_wrap   # AI Assessment (col G)
    new_row[7].alignment = left_wrap   # AI Answer (col H)

    # Hyperlink: Image
    if img_path:
        new_row[2].hyperlink = img_path
        new_row[2].font = Font(color="0000FF", underline="single")
    # Hyperlink: Video Used → clickable link to local video file
    if video_used and video_used != "N/A":
        vid_filename = video_used.split(" [")[0].strip() if " [" in video_used else video_used
        video_local_path = os.path.join(VIDEO_DIR, vid_filename)
        new_row[3].hyperlink = video_local_path
        new_row[3].font = Font(color="0000FF", underline="single")
    # Hyperlink: Audio files
    if audio_in_path:
        new_row[8].hyperlink = audio_in_path
        new_row[8].font = Font(color="0000FF", underline="single")
    if audio_out_path:
        new_row[9].hyperlink = audio_out_path
        new_row[9].font = Font(color="0000FF", underline="single")
    wb.save(EXCEL_FILE)
    wb.close()

# ==============================================================================
#  SESSION FOLDER HELPER
# ==============================================================================
def create_session_folder():
    now = datetime.now()
    date_folder = now.strftime("%Y-%m-%d")
    session_name = now.strftime("%Y-%m-%d_%H-%M-%S")
    session_dir = os.path.join(LOG_FOLDER, date_folder, session_name)
    os.makedirs(session_dir, exist_ok=True)
    return session_dir

# ==============================================================================
#  YOLO DETECTION SCAN
# ==============================================================================
def scan_for_objects(ai_text):
    detected_list = []
    text_lower = ai_text.lower()
    for obj in YOLO_CLASSES:
        if obj == "monitor":
            if "monitor" in text_lower and "tv" not in detected_list:
                detected_list.append("tv")
        elif obj == "cell phone":
            if "cell phone" in text_lower or "phone" in text_lower:
                detected_list.append("cell phone")
        elif obj in text_lower:
            detected_list.append(obj)
    return ",".join(list(set(detected_list)))

# ==============================================================================
#  TEXT CLEANUP
# ==============================================================================
def clean_for_speech(text):
    text = re.sub(r'<[^>]+>', '', text)
    text = re.sub(r'\*{1,3}(.*?)\*{1,3}', r'\1', text)
    text = re.sub(r'_{1,3}(.*?)_{1,3}', r'\1', text)
    text = re.sub(r'#{1,6}\s*', '', text)
    text = re.sub(r'^[\s]*[-•●▪]\s*', '', text, flags=re.MULTILINE)
    text = re.sub(r'^\s*\d+[.)]\s*', '', text, flags=re.MULTILINE)
    emoji_pattern = re.compile(
        "["
        "\U0001F600-\U0001F64F\U0001F300-\U0001F5FF\U0001F680-\U0001F6FF"
        "\U0001F1E0-\U0001F1FF\U00002702-\U000027B0\U000024C2-\U0001F251"
        "\U0001f926-\U0001f937\U00010000-\U0010ffff"
        "\u200d\u2640-\u2642\u2600-\u2B55\u23cf\u23e9\u231a\ufe0f"
        "]+", flags=re.UNICODE
    )
    text = emoji_pattern.sub('', text)
    text = re.sub(r'\([^)]*\)', '', text)
    text = re.sub(r'`+', '', text)
    text = re.sub(r'\n+', ' ', text)
    text = re.sub(r'\s{2,}', ' ', text)
    return text.strip()

# ==============================================================================
#  VIDEO TRANSCRIPT LOOKUP
# ==============================================================================
def get_video_transcript(video_id):
    """Look up a video's transcript from transcripts.json."""
    if video_id in video_transcripts:
        return video_transcripts[video_id]
    clean_id = video_id.replace(".mp4", "")
    if clean_id in video_transcripts:
        return video_transcripts[clean_id]
    return None

def was_video_recently_played():
    """Check if a video was played in the last 5 interactions."""
    if not os.path.exists(EXCEL_FILE):
        return False
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        wb.close()
        for row in rows[-5:]:
            if row and len(row) >= 4 and row[3] and row[3] != "N/A":
                return True
    except Exception:
        pass
    return False

def is_video_summary_request(user_text):
    """
    Smart detection: is the user asking about a previously shown video?
    Two conditions must be met:
    1. A video was recently played (in last 5 conversations)
    2. The user's message references it (directly or indirectly)
    """
    if not was_video_recently_played():
        return False

    lower = user_text.lower()

    # DIRECT references — user explicitly mentions "video"
    direct_video_words = ["video", "clip", "footage"]
    has_direct_ref = any(w in lower for w in direct_video_words)

    # Question/reference words that indicate they want info about something shown
    reference_words = [
        "showed me", "you showed", "you played", "just watched", "just saw",
        "summarize", "summary", "explain", "main point", "key point",
        "important part", "what was that", "tell me about that",
        "what did it say", "what did it show", "go over", "break down",
        "what were the steps", "repeat that", "say that again",
        "what was the", "can you explain"
    ]
    has_reference = any(r in lower for r in reference_words)

    # If they mention "video" at all — almost certainly about the recent video
    if has_direct_ref:
        # But filter out requests for NEW/DIFFERENT videos
        new_video_requests = [
            "show me a video", "find a video", "play a video", "search for a video",
            "is there a video", "do you have a video", "can you find"
        ]
        # Words that signal they want something DIFFERENT, not info about the old one
        new_video_signals = [
            "another video", "different video", "other video", "new video",
            "better video", "another one", "different one", "something else",
            "not helpful", "wasn't helpful", "didn't help", "didn't work",
            "try again", "find another", "show another", "play another"
        ]
        if any(req in lower for req in new_video_requests):
            return False
        if any(sig in lower for sig in new_video_signals):
            return False
        return True

    # If they use reference words without "video" — still likely about the video
    if has_reference:
        return True

    # Indirect: user asks follow-up about the same medical topic without saying "video"
    # e.g. "what's the most important step?" after watching a CPR video
    followup_patterns = [
        "most important", "what should i remember", "key takeaway",
        "what was that about", "can you repeat", "one more time",
        "i missed", "i didn't catch", "i wasn't focused",
        "i was nervous", "go through it again", "the steps"
    ]
    if any(p in lower for p in followup_patterns):
        return True

    return False

# ==============================================================================
#  PHASE 1: AI ASSESSMENT
# ==============================================================================
def ai_assess(user_text, b64_image, memory):
    assessment_prompt = f"""You are a medical AR assistant with camera vision. You can see what the user sees.

TASK: The user said something and you can see their camera view. Analyze BOTH the words AND the image.
Respond with a JSON object only. No explanation, no markdown, no extra text, no thinking tags.

JSON format:
{{
  "intent": "medical" or "vision" or "general" or "conversation",
  "what_i_see": "brief description of what you see in the image",
  "visual_matches_complaint": true or false,
  "severity": "none" or "minor" or "moderate" or "serious" or "not_applicable",
  "needs_video": true or false,
  "search_topic": "specific medical search query if needs_video is true, otherwise empty string",
  "quick_advice": "if minor issue and no video needed, give brief first aid tip, otherwise empty string",
  "video_title": "short 3-6 word topic title for the video display, e.g. 'Nosebleed Treatment Guide' or 'CPR for Adults'"
}}

CRITICAL RULES FOR INTENT:

1. "medical" = ANY health/medicine question. Includes personal complaints ("my hand is bleeding",
   "my back hurts", "I cut myself") AND knowledge questions ("how do I perform CPR", "how to treat a burn").
   -> For personal complaints: CHECK THE IMAGE. If the user says they have an injury, look at
     the camera and verify.
     - If you SEE the injury (bleeding, cut, wound, burn, etc.) in the camera view (including
       on a phone screen, photo, or any visible form) -> set visual_matches_complaint = true
     - If you do NOT see the injury at all in the camera -> set visual_matches_complaint = false
     - For body parts the camera cannot see (back, scalp, bum) or internal symptoms
       (dizziness, headache, chest pain) -> set visual_matches_complaint = true
   -> For knowledge questions ("how do I", "how to", "what should I do if"): set
     visual_matches_complaint = true (image is irrelevant for knowledge questions).
   -> ALWAYS set needs_video = true for medical intent. Provide a specific search_topic.

2. "vision" = User asks what you can see.

3. "general" = Non-medical, non-vision knowledge questions.

4. "conversation" = Casual talk, greetings, personal info.

IMPORTANT:
- ANY question about CPR, bleeding, burns, injuries, first aid, stretches, pain = MEDICAL
- For medical intent: ALWAYS provide search_topic and set needs_video=true
- video_title should be a clean short title like "Wrist Pain Relief" or "CPR for Adults"

CONVERSATION HISTORY:
{memory}"""

    try:
        response = llm_client.chat.completions.create(
            model=LOCAL_MODEL_NAME,
            messages=[
                {"role": "system", "content": assessment_prompt},
                {"role": "user", "content": [
                    {"type": "text", "text": user_text},
                    {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64_image}"}}
                ]}
            ],
            max_tokens=350,
            temperature=0.1
        )
        raw = response.choices[0].message.content.strip()
        raw = re.sub(r'<[^>]+>', '', raw).strip()
        raw = re.sub(r'^```json\s*', '', raw).strip()
        raw = re.sub(r'^```\s*', '', raw).strip()
        raw = re.sub(r'```$', '', raw).strip()
        assessment = json.loads(raw)
        return assessment
    except json.JSONDecodeError as e:
        print(f"   ⚠️ Assessment JSON parse failed: {e}")
        return {
            "intent": "general", "what_i_see": "unable to assess",
            "visual_matches_complaint": True, "severity": "not_applicable",
            "needs_video": False, "search_topic": "", "quick_advice": "", "video_title": ""
        }
    except Exception as e:
        print(f"   ⚠️ Assessment error: {e}")
        return {
            "intent": "general", "what_i_see": "unable to assess",
            "visual_matches_complaint": True, "severity": "not_applicable",
            "needs_video": False, "search_topic": "", "quick_advice": "", "video_title": ""
        }

# ==============================================================================
#  PHASE 3: AI FINAL RESPONSE
# ==============================================================================
SPEECH_RULES = (
    "STRICT RULES: Reply in plain spoken English only. "
    "NEVER use emojis, emoticons, special symbols, asterisks, bold, italic, markdown, "
    "bullet points, numbered lists, dashes, hashtags, stage directions, or thinking tags. "
    "Respond as if speaking out loud to a real person. "
    "Keep responses under 3 sentences unless asked for detail."
)

def ai_respond(user_text, b64_image, assessment, search_result, memory, transcript_context=""):
    intent = assessment.get("intent", "general")
    visual_match = assessment.get("visual_matches_complaint", True)
    severity = assessment.get("severity", "not_applicable")
    what_i_see = assessment.get("what_i_see", "")
    quick_advice = assessment.get("quick_advice", "")
    is_medical = intent == "medical"

    # Special: video transcript summary
    if transcript_context:
        system_prompt = (
            f"You are a helpful Medical AR assistant. {SPEECH_RULES} "
            f"The user is asking about a video that was previously shown to them. "
            f"Here is the EXACT transcript of that video:\n\n{transcript_context[:3000]}\n\n"
            f"CRITICAL: Answer ONLY based on what is actually said in this transcript. "
            f"Do NOT add information that is not in the transcript. "
            f"Do NOT make up times, durations, or steps that are not mentioned. "
            f"If the transcript covers multiple topics, focus on the part most relevant to what the user is asking about. "
            f"Summarize the key points from the transcript naturally as if explaining to a friend. "
            f"HISTORY:\n{memory}"
        )
    elif intent == "medical" and not visual_match:
        # User claims injury but camera shows nothing → tell them honestly, no video
        system_prompt = (
            f"You are a helpful Medical AR assistant. {SPEECH_RULES} "
            f"The user described an injury but you do NOT see it in the camera. "
            f"You see: '{what_i_see}'. "
            f"Be honest and caring. Tell them you cannot see the area they mentioned. "
            f"Ask them to show the affected area more clearly so you can help them. "
            f"DO NOT load a video yet. DO NOT describe the desk or computer setup. "
            f"Keep response under 2 sentences. "
            f"HISTORY:\n{memory}"
        )
    elif intent == "medical" and visual_match and search_result:
        # Either user showed the injury, OR it's a knowledge question → advice + video
        system_prompt = (
            f"You are a helpful Medical AR assistant. {SPEECH_RULES} "
            f"The user has a medical question or concern. "
            f"You see: '{what_i_see}'. "
            f"Give a brief helpful reply about their situation, then tell them you are "
            f"loading a video titled '{search_result['matched_question']}' that will walk them through it. "
            f"DO NOT ask them to show anything. DO NOT describe unrelated items like the desk. "
            f"Keep response under 3 sentences. "
            f"HISTORY:\n{memory}"
        )
    elif intent == "medical" and not search_result:
        # Medical but no video found — give verbal advice
        system_prompt = (
            f"You are a helpful Medical AR assistant. {SPEECH_RULES} "
            f"The user has a medical question. You see: '{what_i_see}'. "
            f"Answer their question with your best medical knowledge. Be clear and helpful. "
            f"Keep response under 3 sentences. "
            f"HISTORY:\n{memory}"
        )
    elif intent == "vision":
        system_prompt = (
            f"You are a helpful AR assistant with camera vision. {SPEECH_RULES} "
            f"Describe what you see in the image naturally. "
            f"HISTORY:\n{memory}"
        )
    else:
        system_prompt = (
            f"You are a helpful AR assistant. {SPEECH_RULES} "
            f"HISTORY:\n{memory}"
        )

    response = llm_client.chat.completions.create(
        model=LOCAL_MODEL_NAME,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": [
                {"type": "text", "text": user_text},
                {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64_image}"}}
            ]}
        ],
        max_tokens=250
    )
    return clean_for_speech(response.choices[0].message.content)

# ==============================================================================
#  INIT
# ==============================================================================
init_excel()
os.makedirs(LOG_FOLDER, exist_ok=True)

# ==============================================================================
#  ENDPOINTS
# ==============================================================================
@app.post("/ask")
async def ask(request: Request, image: UploadFile = File(...), audio: UploadFile = File(...)):
    print("\n" + "=" * 50)
    print("🚀 NEW MISSION REQUEST RECEIVED")

    session_dir = create_session_folder()
    timestamp = datetime.now().strftime("%H-%M-%S")

    img_path = os.path.join(session_dir, f"view_{timestamp}.jpg")
    audio_in_path = os.path.join(session_dir, f"user_{timestamp}.wav")
    audio_out_path = os.path.join(session_dir, f"ai_{timestamp}.wav")

    with open(img_path, "wb") as f:
        shutil.copyfileobj(image.file, f)
    with open(audio_in_path, "wb") as f:
        shutil.copyfileobj(audio.file, f)

    try:
        # ==============================================================
        #  1. HEAR
        # ==============================================================
        print("   👂 Listening...")
        segments, _ = whisper_model.transcribe(audio_in_path, beam_size=5)
        user_text = " ".join([segment.text for segment in segments]).strip()
        print(f"   📝 User: \"{user_text}\"")

        # Encode image
        with open(img_path, "rb") as image_file:
            b64_image = base64.b64encode(image_file.read()).decode('utf-8')

        # Get memory
        full_history = get_full_memory(user_text)

        # ==============================================================
        #  CHECK: Is user asking about a previously shown video?
        # ==============================================================
        transcript_context = ""
        if is_video_summary_request(user_text):
            last_vid = get_last_video_id()
            if last_vid:
                transcript = get_video_transcript(last_vid)
                if transcript:
                    transcript_context = transcript
                    print(f"   📜 Transcript loaded for video: {last_vid} ({len(transcript)} chars)")
                else:
                    print(f"   📜 No transcript found for: {last_vid}")
            else:
                print("   📜 No previous video found in log")

        # If we have transcript context, skip assessment and go straight to response
        if transcript_context:
            print("   🧠 Video summary mode — skipping assessment")
            ai_text = ai_respond(user_text, b64_image, {"intent": "conversation"}, None, full_history, transcript_context)
            print(f"   🤖 AI: \"{ai_text[:80]}...\"")

            # Speak
            print("   👄 Speaking...")
            audio_data, sample_rate = kokoro.create(ai_text, voice=VOICE_NAME, speed=1.0, lang="en-us")
            sf.write(audio_out_path, audio_data, sample_rate)

            # Log
            append_to_log(
                question=user_text,
                assessment="Video Summary Request",
                answer=ai_text,
                img_path=img_path, video_used="", detected_items="",
                audio_in_path=audio_in_path, audio_out_path=audio_out_path
            )
            print("   📝 Logged")

            safe_user = user_text.encode('ascii', 'ignore').decode()[:500].replace("\n", " ")
            safe_ai = ai_text.encode('ascii', 'ignore').decode()[:500].replace("\n", " ")

            return FileResponse(
                audio_out_path, media_type="audio/wav",
                headers={
                    "Access-Control-Expose-Headers": "X-Detected-Object, X-User-Transcript, X-AI-Response, X-Video-URL, X-Video-Start, X-Video-End, X-Video-Title",
                    "X-Detected-Object": "none",
                    "X-User-Transcript": safe_user,
                    "X-AI-Response": safe_ai,
                    "X-Video-URL": "", "X-Video-Start": "", "X-Video-End": "", "X-Video-Title": ""
                }
            )

        # ==============================================================
        #  2. ASSESS
        # ==============================================================
        print("   🧠 Phase 1: Assessing...")
        assessment = ai_assess(user_text, b64_image, full_history)

        intent = assessment.get("intent", "general")
        visual_match = assessment.get("visual_matches_complaint", True)
        severity = assessment.get("severity", "not_applicable")
        needs_video = assessment.get("needs_video", False)
        search_topic = assessment.get("search_topic", "")
        what_i_see = assessment.get("what_i_see", "")
        video_title = assessment.get("video_title", "")

        print(f"   📋 Assessment:")
        print(f"      Intent: {intent}")
        print(f"      I see: {what_i_see}")
        print(f"      Visual matches: {visual_match}")
        print(f"      Severity: {severity}")
        print(f"      Needs video: {needs_video}")
        if search_topic:
            print(f"      Search for: \"{search_topic}\"")
        if video_title:
            print(f"      Video title: \"{video_title}\"")

        # ==============================================================
        #  3. SEARCH — Smart decision based on intent + visual match
        #
                                # ==============================================================
        video_url = ""
        video_start = ""
        video_end = ""
        video_used = ""
        search_result = None

        should_search = False
        if needs_video and search_topic and intent == "medical":
            if not visual_match:
                # User claims injury but camera shows nothing → don't search
                print("   📹 Injury not visible in camera — no search, asking user to show area")
                should_search = False
            else:
                # Visual match OR knowledge question → search for video
                should_search = True

        if should_search and search_topic:
            print(f"   🔍 Phase 2: Searching \"{search_topic}\"")
            search_result = search_engine.search(search_topic, use_llm=True)

            if search_result:
                video_url = search_result['video_file']
                video_start = str(search_result['start_second'])
                video_end = str(search_result['end_second'])
                video_used = f"{search_result['video_file']} [{video_start}s-{video_end}s]"
                print(f"   📹 Found: {video_used}")
            else:
                # Fallback: try raw user text
                print(f"   🔍 Fallback: \"{user_text}\"")
                search_result = search_engine.search(user_text, use_llm=True)
                if search_result:
                    video_url = search_result['video_file']
                    video_start = str(search_result['start_second'])
                    video_end = str(search_result['end_second'])
                    video_used = f"{search_result['video_file']} [{video_start}s-{video_end}s]"
                    print(f"   📹 Fallback found: {video_used}")
                else:
                    print("   📹 No video found")
        elif intent == "medical":
            print("   📹 No search needed")
        else:
            print("   📹 Not medical — no search")

        # ==============================================================
        #  4. RESPOND
        # ==============================================================
        print("   🧠 Phase 3: Responding...")
        ai_text = ai_respond(user_text, b64_image, assessment, search_result, full_history)
        print(f"   🤖 AI: \"{ai_text[:80]}...\"")

        # ==============================================================
        #  5. DETECT
        # ==============================================================
        detected_keyword = ""
        if intent == "vision":
            detected_keyword = scan_for_objects(ai_text)
            if detected_keyword:
                print(f"   🎯 Detected: {detected_keyword}")
        else:
            print("   🎯 Detection skipped")

        # ==============================================================
        #  6. SPEAK
        # ==============================================================
        print("   👄 Speaking...")
        audio_data, sample_rate = kokoro.create(ai_text, voice=VOICE_NAME, speed=1.0, lang="en-us")
        sf.write(audio_out_path, audio_data, sample_rate)

        # ==============================================================
        #  7. LOG (formatted assessment)
        # ==============================================================
        assessment_log = (
            f"Intent: {intent}\n"
            f"See: {what_i_see[:60]}\n"
            f"Match: {visual_match} | Severity: {severity}\n"
            f"Video: {video_used if video_used else 'N/A'}"
        )
        append_to_log(
            question=user_text,
            assessment=assessment_log,
            answer=ai_text,
            img_path=img_path,
            video_used=video_used,
            detected_items=detected_keyword,
            audio_in_path=audio_in_path,
            audio_out_path=audio_out_path
        )
        print("   📝 Logged")

        # ==============================================================
        #  8. RETURN
        # ==============================================================
        safe_user = user_text.encode('ascii', 'ignore').decode()[:500].replace("\n", " ")
        safe_ai = ai_text.encode('ascii', 'ignore').decode()[:500].replace("\n", " ")
        safe_title = video_title.encode('ascii', 'ignore').decode()[:100] if video_title else ""
        header_detect = detected_keyword if detected_keyword else "none"

        print(f"   📤 → Detect: {header_detect} | Video: {video_url} | Title: {safe_title}")

        return FileResponse(
            audio_out_path,
            media_type="audio/wav",
            headers={
                "Access-Control-Expose-Headers": "X-Detected-Object, X-User-Transcript, X-AI-Response, X-Video-URL, X-Video-Start, X-Video-End, X-Video-Title",
                "X-Detected-Object": str(header_detect),
                "X-User-Transcript": str(safe_user),
                "X-AI-Response": str(safe_ai),
                "X-Video-URL": video_url,
                "X-Video-Start": video_start,
                "X-Video-End": video_end,
                "X-Video-Title": safe_title
            }
        )

    except Exception as e:
        print(f"   ❌ ERROR: {e}")
        import traceback
        traceback.print_exc()
        return FileResponse(audio_in_path, media_type="audio/wav")


@app.post("/log_screenshot")
async def log_screenshot(file: UploadFile = File(...)):
    now = datetime.now()
    date_folder = now.strftime("%Y-%m-%d")
    ar_dir = os.path.join(LOG_FOLDER, date_folder, "AR_Screenshots")
    os.makedirs(ar_dir, exist_ok=True)
    filename = f"detection_AR_{now.strftime('%H-%M-%S')}.jpg"
    file_location = os.path.join(ar_dir, filename)
    with open(file_location, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    return {"status": "saved", "path": file_location}


if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=8000)
